"""Build the Lab CA report from measured experiment artifacts.

Run with Python containing python-docx, reportlab and Pillow. Rendering QA is a
separate step so the PDF remains reproducible without a desktop office suite.
"""
from __future__ import annotations

import csv
import json
import textwrap
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image as PILImage, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate, Frame, Image, KeepTogether, PageBreak, PageTemplate,
    Paragraph, Preformatted, Spacer, Table, TableStyle,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "report"
ASSETS = OUT / "generated"
ART = ROOT / "artifacts"
SOURCE = "Local source folder: RL_LABCA"
GAME = "TicTacToe_Offline.html"
NOTEBOOK = "notebooks/TicTacToe_Q_Learning.ipynb"
FONTDIR = Path("C:/Windows/Fonts")


def read_json(path, default=None):
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else default


def h(text, level=1):
    return {"type": "heading", "text": text, "level": level}


def p(text):
    return {"type": "paragraph", "text": text}


def bullets(items):
    return {"type": "bullets", "items": items}


def table(headers, rows, widths=None):
    return {"type": "table", "headers": headers, "rows": rows, "widths": widths}


def pic(path, caption, width=6.5, max_height=3.8):
    return {"type": "image", "path": str(path), "caption": caption, "width": width, "max_height": max_height}


def code(text):
    return {"type": "code", "text": textwrap.dedent(text).strip()}


def font(size, bold=False):
    return ImageFont.truetype(str(FONTDIR / ("arialbd.ttf" if bold else "arial.ttf")), size)


def create_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    im = PILImage.new("RGB", (1800, 780), "white")
    d = ImageDraw.Draw(im)
    boxes = [(60, 95, 530, 285), (665, 95, 1135, 285), (1270, 95, 1740, 285),
             (1270, 455, 1740, 645), (665, 455, 1135, 645), (60, 455, 530, 645)]
    labels = [("Environment", "Board and opponent"), ("State", "Agent perspective and symmetry"),
              ("RL agent", "Legal epsilon greedy action"), ("Action and reward", "Agent move then opponent move"),
              ("State update", "Next agent decision or terminal"), ("Learning", "One step Q table update")]
    for box, (title, subtitle) in zip(boxes, labels):
        d.rounded_rectangle(box, radius=16, fill="#eef2f6", outline="#64748b", width=3)
        x = (box[0] + box[2]) // 2
        d.text((x, box[1] + 55), title, font=font(36, True), fill="black", anchor="mm")
        for idx, line in enumerate(textwrap.wrap(subtitle, 29)):
            d.text((x, box[1] + 108 + idx * 34), line, font=font(27), fill="#253246", anchor="mm")
    def arrow(points, direction):
        d.line(points, fill="#334155", width=6)
        x, y = points[-1]
        tri = {"r": [(x,y),(x-21,y-13),(x-21,y+13)], "l": [(x,y),(x+21,y-13),(x+21,y+13)],
               "d": [(x,y),(x-13,y-21),(x+13,y-21)], "u": [(x,y),(x-13,y+21),(x+13,y+21)]}[direction]
        d.polygon(tri, fill="#334155")
    arrow([(530,190),(665,190)], "r")
    arrow([(1135,190),(1270,190)], "r")
    arrow([(1505,285),(1505,455)], "d")
    arrow([(1270,550),(1135,550)], "l")
    arrow([(665,550),(530,550)], "l")
    arrow([(295,455),(295,285)], "u")
    d.text((900,733), "Repeat until win, loss or draw; reset the environment for the next episode", font=font(29), fill="#475569", anchor="mm")
    im.save(ASSETS / "workflow.png")

    # A rendered, structured equation avoids ambiguous ASCII math in the report.
    eq = PILImage.new("RGB", (2100, 250), "white")
    e = ImageDraw.Draw(eq)
    terms = [("Q(s,a) ← Q(s,a) + α [ r + γ ", 56), ("max", 56), (" Q(s′,a′) − Q(s,a) ]", 56)]
    x = 55
    for idx, (t, size) in enumerate(terms):
        e.text((x,85), t, font=font(size), fill="black")
        width = e.textlength(t, font=font(size))
        if idx == 1:
            e.text((x+width/2,160), "a′ ∈ A(s′)", font=font(31), fill="black", anchor="mt")
        x += width
    eq.save(ASSETS / "q_update.png")


def find_picture(names):
    for name in names:
        path = ART / name
        if path.exists():
            return path
    return None


def build_pages(context):
    """Content schema shared by Markdown, DOCX and PDF writers."""
    c = context
    pages = []
    pages.append([
        h("Reinforcement Learning Lab CA", 0),
        p("Mini Project Report"),
        table(["Course Name", "Reinforcement Learning", "Semester", "VII"], [
            ["Date of Performance", "___ / ___ / ______", "DIV / Batch No", "________________"],
            ["Student Name", "________________________", "Roll No", "________________"],
        ], [1.25, 2.45, 1.1, 1.7]),
        h("1. Title"),
        p("Tic Tac Toe Agent Using Tabular Q Learning with an Interactive Game Interface"),
        h("Project overview", 2),
        p("This project builds a reinforcement learning agent for the two player game Tic Tac Toe. The agent learns state-action values from repeated games and uses the resulting Q table to select legal moves. A fully offline browser interface lets a human play against the trained policy, inspect action values and restart games. A local Jupyter notebook exposes the training, evaluation and demonstration workflow and remains compatible with Google Colab."),
        p("The experiment separates learning from evaluation. A randomly acting baseline and the trained greedy policy are evaluated as both X and O against random, tactical and optimal minimax opponents. Wins, draws, losses, terminal reward and game length provide complementary views of policy quality. Evaluation uses independent random seeds and no exploration."),
        p(c["overview_result"]),
        h("Submission contents", 2),
        bullets(["Python environment, Q-learning agent, training and evaluation scripts, and reproducible experiment artifacts.",
                 "A self-contained offline browser game with its learned Q table embedded, and a locally runnable notebook.",
                 "Measured results, learning curves, working screenshots, implementation explanations and a viva preparation section."]),
        p("The report covers the problem, RL formulation, implementation, measured performance and limitations. References and source code excerpts follow the main report."),
    ])
    pages.append([
        h("2. Problem Definition and RL Formulation"),
        h("2.1 Problem Statement", 2),
        p("The task is to learn a decision policy that plays Tic Tac Toe effectively without storing a hand-written rule for every possible board. Two players alternate marking empty cells on a 3 by 3 board. X moves first. A player wins by completing a row, column or diagonal; a full board without a winning line is a draw. The agent must choose only legal actions and should maximize its expected discounted return over a game."),
        p("Tic Tac Toe is a controlled simulator for sequential decision making rather than a deployment in a physical system. It makes delayed consequences visible: a locally attractive move can allow an opponent to form a fork several turns later. The complete board is observable, actions are discrete, and each game ends within nine moves. These properties make the problem small enough for a tabular learning method and for direct inspection of the learned values."),
        h("2.2 Motivation", 2),
        p("A supervised classifier would require labeled best moves. Q-learning instead uses interaction and reward. The agent explores moves, observes whether they eventually lead to a win, draw or loss, and adjusts the value of earlier choices. The project therefore demonstrates exploration, delayed reward, temporal difference updates and policy evaluation in an environment whose rules can be verified independently."),
        p("The game is also useful for understanding the limits of reinforcement learning. Strong play against a weak opponent does not imply perfect play against an optimal opponent. Evaluation against several opponent types, with separate results for starting and second player roles, exposes this distinction. The interface connects the numerical Q table with an observable sequence of decisions so a learner can explain both the update rule and its practical effect."),
        h("2.3 Objectives", 2),
        bullets(["Implement a correct Tic Tac Toe environment with legal action masking, terminal detection and reproducible experiments.",
                 "Train a tabular Q-learning policy through trial and error and save a reusable Q table.",
                 "Compare the untrained baseline and trained policy across opponent strategies and player roles using measured metrics and learning curves.",
                 "Provide a usable browser game and a Colab notebook that demonstrate the agent and expose its training process."]),
        h("Scope of the project", 2),
        p("The project uses the standard 3 by 3 game. The agent observes the full board; no image recognition or external dataset is required. The opponent is part of the environment. Minimax is used as an opponent and evaluation reference, while the trained agent itself selects actions from its learned Q table."),
    ])
    pages.append([
        h("2.4 RL Problem Formulation"),
        table(["RL component", "Description"], [
            ["Environment", "A standard 3 by 3 Tic Tac Toe board together with the opponent policy, legal move rules and terminal outcome detector."],
            ["State S", "Nine cells encoded from the learner's perspective: own mark, opponent mark or empty. Only states at an agent decision are learned. Board rotations and reflections share a canonical state internally."],
            ["Action A", "Select one of the empty cells, indexed 0 to 8 in row-major order. Occupied cells are masked from exploration, greedy selection and bootstrap maximization."],
            ["Reward R", c["reward_description"]],
            ["Policy π", "During training, epsilon-greedy action selection over legal cells. During evaluation and normal trained play, choose a legal action with maximal learned Q value."],
            ["Discount factor γ", c["gamma"] + "; discounts return between successive agent decisions."],
            ["Episode and termination", "One complete game. The episode ends immediately after a player wins or the board fills without a winner. A reset clears the board for the next game."],
        ], [1.35,5.15]),
        h("A transition includes the opponent response", 2),
        p("An agent action does not usually lead directly to its next decision. The environment first applies the agent's move. If that move does not end the game, it applies the opponent's response. The resulting board is the next state for Q-learning. Thus one learning transition normally spans two individual moves. A terminal move may shorten the transition. Bootstrapping from a board where it is the opponent's turn would assign values to the wrong decision maker and is avoided."),
        h("State normalization and symmetry", 2),
        p("Encoding marks relative to the learner lets one table support both X and O. The eight square symmetries represent strategically equivalent boards. Canonicalization transforms the state and its action indices together; a selected action is mapped back to the original board before it is played. The web export expands learned values into ordinary board orientation so the browser can use a direct state lookup."),
        p("The opponent mixture remains fixed during training. This produces a stationary stochastic environment at the agent's decision times. Changing the opponent behavior after training changes the distribution of encountered states and can reveal weaknesses that were uncommon in training."),
    ])
    pages.append([
        h("2.5 Selected RL Algorithm"),
        p("Name of algorithm: tabular Q-learning with an epsilon-greedy behavior policy."),
        p("Q-learning estimates the expected return of choosing an action in a state and then following a greedy policy. It is a model-free, off-policy temporal difference method: the behavior policy may explore, while the update target uses the highest value among legal next actions. The Tic Tac Toe state space is small enough to store these estimates explicitly, making the method easier to inspect than a neural value function."),
        h("Q value update", 2),
        pic(ASSETS / "q_update.png", "Equation 1. One step Q-learning update for a nonterminal transition.", max_height=1.0),
        p("Here s is the current state, a is the chosen action, r is the observed reward, s′ is the next agent decision state, α is the learning rate and γ is the discount factor. The maximization considers only legal actions at s′. For a terminal transition, the bootstrap term is zero and the target is just the terminal reward. The temporal difference error is the target minus the current estimate."),
        h("Exploration and exploitation", 2),
        p("With probability epsilon, the agent samples a legal move to explore. Otherwise, it chooses a legal move with the largest current Q value. Epsilon decreases across training so early games cover alternatives and later games make more use of learned values. Random selection among tied legal maxima avoids introducing a permanent cell-order preference during learning and evaluation."),
        h("Reason for selecting the algorithm", 2),
        bullets(["The state and action spaces are discrete and small, so a dictionary-based Q table is practical.",
                 "The update can be shown directly in code and explained using a single observed transition.",
                 "No model of future game outcomes or labeled training dataset is required by the learner.",
                 "The saved policy can be exported to a small client-side game without a model-serving service."]),
        h("Interpretation of the learned values", 2),
        p("A Q value is an estimated discounted return under the training environment. It is not a calibrated win probability. The result depends on reward choices, discounting, opponent behavior, learning rate and state coverage. The browser's action-value view is therefore useful for comparison between legal actions on a board, but it should not be read as a guaranteed outcome against every opponent."),
        p("Classical convergence results require specific conditions, including sufficient repeated exploration and suitable learning-rate schedules. This finite experiment uses a practical fixed learning rate, so stable measured performance is treated as empirical evidence rather than proof of convergence to an optimal policy [1, 2]."),
    ])
    pages.append([
        h("3. Implementation and Working of RL Agent"),
        h("3.1 System Architecture and Workflow", 2),
        pic(ASSETS / "workflow.png", "Figure 1. Learning loop at the agent's decision times.", max_height=2.8),
        p("The environment owns game rules and board transitions. The opponent policies provide reproducible behavior at several strengths. The agent stores and updates state-action values. Training records a history and saves a learned policy; evaluation loads that policy without updating it. A single HTML file embeds the interface and exported policy, allowing browser play without a server or internet connection. The local notebook runs the Python workflow and displays its outputs."),
        h("3.2 Tools and Technologies", 2),
        table(["Category", "Project technology and use"], [
            ["Programming languages", "Python for the simulator, Q-learning, training and evaluation; JavaScript, HTML and CSS for the browser game."],
            ["Libraries", c["libraries"]],
            ["Development environment", "Local Python execution and Jupyter notebook; source files retained in the project folder. The notebook can also be opened in Google Colab."],
            ["Simulator and dataset", "A custom Tic Tac Toe simulator generates training trajectories. No external training dataset is used."],
            ["Saved artifacts", "Q table in JSON, evaluation and training summaries in JSON, training history in CSV, and plotted learning curves."],
            ["Game delivery", "TicTacToe_Offline.html bundles HTML, CSS, JavaScript and the learned Q table. Double-click to play locally with no backend or network dependency."],
        ], [1.35,5.15]),
        p("The separation between game rules, learning and presentation makes each component easier to inspect. Replaying the learned table does not require retraining. The training scripts can regenerate the experiment artifacts, and the notebook presents the same process in an interactive environment."),
    ])
    pages.append([
        h("3.3 Algorithm Implementation"),
        h("Environment and legal actions", 2),
        p("The board stores nine cell values. A legal-action function returns the indices of empty cells. Terminal detection checks the eight possible winning lines before declaring a full-board draw. Applying a move validates its legality, updates the board and determines whether the episode has ended. These rules are shared across training and evaluation so that metrics describe the same game."),
        h("Training sequence", 2),
        bullets(["Reset the board and assign the learner to X or O. If the learner is O, allow the opponent to make the opening move.",
                 "Convert the current board to the learner's perspective and canonical orientation.",
                 "Choose a legal action with the current epsilon-greedy policy and map the action back to the board.",
                 "Apply the learner move. If the game continues, apply the opponent response and inspect the outcome again.",
                 "Compute reward and the next state. Use reward alone at a terminal state; otherwise add the discounted best legal next-state value.",
                 "Update the selected Q entry, record outcomes and continue until the episode terminates."]),
        h("Opponent strategies", 2),
        table(["Opponent", "Behavior and purpose"], [
            ["Random", "Uniformly samples legal moves. Provides varied trajectories and a low-strength comparison."],
            ["Tactical", c["tactical_description"]],
            ["Minimax", "Searches the remaining legal game tree to choose an optimal outcome. Serves as a strong training component and an evaluation reference."],
        ], [1.15,5.35]),
        h("Policy export and interface integration", 2),
        p("After training, the Q table and training metadata are written to JSON. Symmetric states are expanded into the orientation expected by the browser. At each computer turn, the interface forms the same perspective state, masks occupied cells, and selects a legal action from the loaded values. User controls expose player role and game state; the learned values remain available for explanation."),
        p("The user-facing trained mode plays from Q values. The optimal search opponent is kept distinct in experiments so measured Q-learning behavior is not silently replaced with a search solution."),
    ])
    pages.append([
        h("3.4 Hyperparameters and Parameter Settings"),
        table(["Parameter", "Value"], c["hyper_rows"], [2.35,4.15]),
        h("3.5 Training Procedure", 2),
        p(c["training_description"]),
        p("Each game supplies a short sequence of decisions. Intermediate nonterminal transitions receive the configured step reward and use a discounted bootstrap target. Winning, losing or filling the board ends the episode, so terminal updates never read future action values. Rewards obtained after an opponent reply are attributed to the learner's preceding action through the same transition."),
        p("The learning rate controls how strongly a new target changes an existing value. A constant learning rate continues to adapt late in training, while epsilon decay shifts the behavior policy toward exploitation. Alternating player roles prevents training from covering only first-player openings. Symmetry reduces duplicated experience, but action indices must be transformed with the board for the update to remain correct."),
        h("Reproducibility and evaluation separation", 2),
        p(c["evaluation_protocol"]),
        p("Training histories summarize the behavior actually used during learning, including exploration. Checkpoint evaluations use a greedy policy on separate games. These answer different questions: an exploratory training curve describes experience collection, whereas an evaluation curve estimates the current policy's performance without exploratory moves. The report labels them separately."),
    ])
    demo = [
        h("3.6 Working Demonstration"),
        p("The browser demonstration loads the trained Q table and accepts human moves on the board. The game reports whose turn it is, detects terminal outcomes, prevents moves in occupied cells and supports restarting. The interface exposes the learner's values so a selected move can be related to the Q-learning policy."),
    ]
    for imgpath, caption in c["screenshots"][:2]:
        demo.append(pic(imgpath, caption, max_height=2.65))
    demo.extend([
        h("Demonstration procedure", 2),
        bullets(["Double-click TicTacToe_Offline.html and start a new match with the desired player role.",
                 "Select an empty cell and observe the learned agent's reply and action values.",
                 "Continue until the game reports a win, loss or draw, then restart and change roles.",
                 "Open the local notebook, run its cells and inspect the training and evaluation outputs."]),
        p(c["execution_status"]),
    ])
    pages.append(demo)
    pages.append([
        h("4. Results Analysis and Viva"),
        h("4.1 Experimental Results", 2),
        p(c["results_intro"]),
        table(["Policy", "Opponent", "Games", "Win %", "Draw %", "Loss %", "Avg reward"], c["result_rows"], [.85,1.05,.58,.6,.62,.6,1.2]),
        p(c["results_interpretation"]),
        h("Performance by player role", 2),
        table(["Role", "Opponent", "Games", "Wins", "Draws", "Losses", "Non-loss %"], c["role_rows"], [.55,1.25,.65,.65,.65,.65,1.1]),
        p("X has the opening move and O responds to the opening, so pooled performance can conceal role-specific weaknesses. These separate results help distinguish broad policy strength from a policy that mainly succeeds when moving first."),
        h("Reading the result tables", 2),
        p("Win rate and non-loss rate answer different questions. A draw against an optimal opponent is a successful defensive outcome, while a high draw rate against a random opponent may indicate missed winning opportunities. The loss rate remains essential even when average reward appears favorable."),
    ])
    graph_page = [
        h("4.2 Performance Metrics"),
        table(["Metric", "Definition and interpretation"], [
            ["Cumulative reward", "Sum of rewards over evaluated episodes. Its scale depends on the game count and configured terminal rewards."],
            ["Average reward", "Total episode reward divided by games. Reported alongside outcome rates because reward depends on the chosen draw value."],
            ["Win rate", "Wins divided by games. Measures conversion of games into victories."],
            ["Non-loss rate", "Wins plus draws divided by games. Useful against strong opponents where drawing can be the best achievable result."],
            ["Game length", "Individual board moves per episode, not the number of Q updates. Terminal games contain at most nine moves."],
            ["Empirical stability", "Whether measured checkpoint performance stabilizes. It is not a mathematical proof of convergence."],
        ], [1.25,5.25]),
        h("4.3 Graphs and Visualization", 2),
    ]
    if c.get("reward_plot"):
        graph_page.append(pic(c["reward_plot"], c["reward_caption"], max_height=3.8))
    graph_page.append(p(c["reward_analysis"]))
    pages.append(graph_page)
    trends = [h("4.3 Learning Curves Continued")]
    if c.get("success_plot"):
        trends.append(pic(c["success_plot"], c["success_caption"], max_height=2.6))
    if c.get("steps_plot"):
        trends.append(pic(c["steps_plot"], c["steps_caption"], max_height=2.6))
    trends.extend([
        p(c["curve_analysis"]),
        h("4.4 Result Analysis", 2),
        p(c["analysis"]),
        p(c["audit_analysis"]),
    ])
    pages.append(trends)
    pages.append([
        h("4.5 Limitations"),
        bullets(["The environment is a fully observable 3 by 3 game. The tabular method does not scale directly to games with large or continuous state spaces.",
                 "Performance depends on the training opponent mixture. A policy can exploit common mistakes while retaining weaknesses against adversarial move sequences.",
                 "The experiment uses a finite training budget and a fixed learning rate. Stable curves do not establish the theoretical convergence conditions of Q-learning.",
                 "Sampled evaluation results describe specified opponents and seeds. The separate exact audit covers this greedy policy from the empty board; it does not certify every arbitrary board setup or exploratory difficulty setting.",
                 "Canonicalization and perspective encoding reduce duplicate states but demand correct state-action mapping. Incorrect transformations would silently corrupt action values.",
                 "Action values are reward estimates, not probabilities. Interface users should interpret their ranking rather than treating a displayed value as a certainty."]),
        h("4.6 Future Scope"),
        p("An immediate extension is an audit of every valid board setup, including positions that the saved greedy policy never reaches from the empty board. This can identify a board where the learned action sacrifices an available draw or win after a forced demonstration move. Training across several independent seeds and reporting confidence intervals would better characterize variation than one reproducible run."),
        p("The project can compare Q-learning with SARSA, Monte Carlo control and dynamic programming under a common evaluation protocol. Additional experiments can vary the learning rate, epsilon schedule, discount factor and opponent mixture. A decaying learning-rate schedule and visitation counts would support a clearer study of convergence conditions. Larger boards or Connect Four would motivate function approximation, replay-based learning and more careful generalization tests."),
        p("The interface can add move-by-move replays and an explanation that links a selected action to its nearest competing legal action. A teaching mode could pause before the computer move, ask the learner to predict the update, and then reveal the reward and temporal difference error."),
        h("4.7 Conclusion"),
        p(c["conclusion"]),
        p("The implementation connects the full RL workflow: a simulator produces experience, Q-learning updates action values, separate experiments measure performance, and an interactive game demonstrates the resulting behavior. The main lesson is that a learned policy must be assessed against the opponents and decision states relevant to its intended use, with wins, draws and losses all reported."),
    ])
    pages.append([
        h("Viva Preparation"),
        h("Why is this reinforcement learning", 2),
        p("The agent learns from its own state-action-reward transitions. No dataset provides a correct move label for every board. The objective is long-term return over a sequence of decisions."),
        h("Why does the transition include an opponent move", 2),
        p("The next Q state must be another point where the same agent can choose an action. Combining the learner move and opponent reply preserves that decision perspective. Terminal outcomes are handled immediately."),
        h("Why is Q-learning off policy", 2),
        p("The behavior policy can choose random exploratory actions, while the target uses the maximum legal next-action value. The target therefore evaluates a greedy continuation rather than the action actually chosen by the exploratory policy."),
        h("Why are illegal moves masked", 2),
        p("Only empty cells are valid. An occupied cell must be excluded from both action selection and the next-state maximum; otherwise an unused or arbitrary Q value could distort the update target."),
        h("Does a high Q value equal a high win probability", 2),
        p("No. Q values include discounting and the numeric rewards assigned to wins, draws and losses. They estimate return under the training environment rather than a calibrated probability."),
        h("Why test against minimax", 2),
        p("Minimax provides an optimal game-playing reference. Performance against random play can hide exploitable mistakes, while optimal opposition tests whether the learned policy can preserve a draw when a win cannot be forced."),
        h("Why report both X and O", 2),
        p("The two roles face different opening conditions. Reporting them separately can reveal whether the policy learned only strong first-player behavior or also learned to respond reliably as the second player."),
        h("Can these curves prove convergence", 2),
        p("No. They show empirical performance for a finite run. Formal Q-learning convergence depends on repeated state-action coverage and appropriate learning-rate conditions, which cannot be concluded from a smooth reward plot."),
        h("How would you explain one update in the viva", 2),
        p("Identify the board, legal action, received reward and next agent decision state. Find the largest legal next-state Q value, compute the target and subtract the current value to obtain the temporal difference error. Multiply that error by the learning rate and add it to the old Q value. Use zero bootstrap at a terminal state."),
    ])
    pages.append([
        h("5. References"),
        p("[1] Richard S. Sutton and Andrew G. Barto. Reinforcement Learning: An Introduction. Second edition. MIT Press, 2018. Used for the reinforcement learning framework, temporal difference learning and off-policy control."),
        p("https://mitpress.mit.edu/9780262039246/reinforcement-learning/"),
        p("[2] Christopher J. C. H. Watkins and Peter Dayan. Q-learning. Machine Learning, 8, 279-292, 1992. Used for the Q-learning algorithm and the distinction between practical finite training and theoretical convergence conditions."),
        p("https://www.gatsby.ucl.ac.uk/~dayan/papers/cjch.pdf"),
        p("[3] Google. Colaboratory Frequently Asked Questions. Used for the notebook execution environment and its hosted runtime model. Accessed 4 October 2026."),
        p("https://research.google.com/colaboratory/intl/en-GB/faq.html"),
        p("[4] Project source code, trained policy and reproducible experiment artifacts. Tic Tac Toe Q Learning Lab. The local project is the primary source for all measured numerical results in this report."),
        p(SOURCE),
        h("Project access", 2),
        p("Source code: tictactoe/ and scripts/ in the project folder."),
        p("Offline browser game: " + GAME),
        p("Local notebook: " + NOTEBOOK),
        h("Experiment evidence", 2),
        p("The machine-readable training configuration and summary are stored in artifacts/training_summary.json. Evaluation outcomes are stored in artifacts/evaluation.json. Training history and plotted figures accompany the report in the same local project. The report builder reads these saved outputs so displayed numerical results remain traceable to the experiment."),
        p(c["provenance"]),
    ])
    pages.append([
        h("Appendix A Source Code and Reproduction"),
        p("The local project folder contains the complete runnable implementation. The following commands show the execution path; the notebook provides an interactive version of training and evaluation. The offline game itself needs only a browser."),
        code(c["run_commands"]),
        h("Q-learning update", 2),
        p("The excerpt below is taken from the project implementation. It shows the terminal target and the legal-action maximum used in the value update."),
        code(c["update_code"]),
        h("Reading the code", 2),
        p("The update changes only the value of the state-action pair that produced the observed transition. A zero bootstrap at termination prevents the terminal result from being diluted by values on a board where no future decision exists. Legal masking ensures that an impossible occupied-cell move cannot become the target."),
        p("To reproduce the report, regenerate training and evaluation artifacts with the stated configuration, then run scripts/build_report.py in an environment with python-docx, reportlab and Pillow. The committed artifacts preserve the reported run even when a notebook user chooses a shorter demonstration budget."),
    ])
    pages.append([
        h("Appendix B Sample Outputs and Verification"),
        p(c["sample_output_intro"]),
        code(c["sample_output"]),
        h("Important implementation checks", 2),
        bullets(c["verification"]),
        h("Files to inspect during demonstration", 2),
        table(["File or folder", "Purpose"], [
            ["tictactoe/", "Game rules, Q-learning behavior, opponent policies and experiment implementation."],
            ["artifacts/q_table.json", "Saved learned values and policy metadata used by the browser."],
            ["artifacts/training_summary.json", "Training configuration and run summary."],
            ["artifacts/evaluation.json", "Measured baseline and trained evaluation outcomes."],
            ["artifacts/training_history.csv", "Recorded learning history used to make graphs."],
            ["notebooks/TicTacToe_Q_Learning.ipynb", "Local notebook for interactive training, evaluation and demonstration; compatible with Google Colab."],
            ["report/", "Submission report in PDF, Word and Markdown formats."],
        ], [2.65,3.85]),
        p("Personal academic fields on the cover are intentionally left for the student to complete before submission."),
    ])
    return pages


def register_fonts():
    for name, filename in [("Arial", "arial.ttf"), ("Arial-Bold", "arialbd.ttf"),
                           ("Arial-Italic", "ariali.ttf"), ("Consolas", "consola.ttf")]:
        pdfmetrics.registerFont(TTFont(name, str(FONTDIR / filename)))
    pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold", italic="Arial-Italic")


def image_size(block):
    with PILImage.open(block["path"]) as im:
        w, hh = im.size
    scale = min(block["width"] * 72 / w, block["max_height"] * 72 / hh)
    return w * scale, hh * scale


def make_pdf(pages):
    register_fonts()
    styles = {
        "body": ParagraphStyle("Body", fontName="Arial", fontSize=10.5, leading=14.3, spaceAfter=8),
        "title": ParagraphStyle("Title", fontName="Arial-Bold", fontSize=24, leading=29, spaceAfter=12),
        "h1": ParagraphStyle("H1", fontName="Arial-Bold", fontSize=17, leading=21, spaceAfter=12, keepWithNext=True),
        "h2": ParagraphStyle("H2", fontName="Arial-Bold", fontSize=12, leading=16, spaceBefore=9, spaceAfter=6, keepWithNext=True),
        "caption": ParagraphStyle("Caption", fontName="Arial-Italic", fontSize=8.7, leading=11.5, spaceBefore=5, spaceAfter=8),
        "cell": ParagraphStyle("Cell", fontName="Arial", fontSize=9.0, leading=12),
        "th": ParagraphStyle("TH", fontName="Arial-Bold", fontSize=9, leading=12, textColor=colors.white),
        "code": ParagraphStyle("Code", fontName="Consolas", fontSize=8.2, leading=11.4, spaceAfter=9),
    }
    out = OUT / "LAB_CA_Report.pdf"
    doc = BaseDocTemplate(str(out), pagesize=(595.276,841.89), leftMargin=48, rightMargin=48,
                          topMargin=43, bottomMargin=42, title="Tic Tac Toe Agent Using Tabular Q Learning",
                          author="", allowSplitting=1)
    frame = Frame(doc.leftMargin,doc.bottomMargin,doc.width,doc.height, leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)
    def footer(canvas, doc):
        canvas.setFont("Arial",8)
        canvas.setFillColor(colors.HexColor("#555555"))
        canvas.drawString(48,24,"Reinforcement Learning Lab CA")
        canvas.drawRightString(547,24,str(doc.page))
    doc.addPageTemplates(PageTemplate(id="report",frames=[frame],onPage=footer))
    story=[]
    for page_index,page in enumerate(pages):
        if page_index:
            story.append(PageBreak())
        for b in page:
            typ=b["type"]
            if typ=="heading":
                style=styles["title" if b["level"]==0 else "h1" if b["level"]==1 else "h2"]
                story.append(Paragraph(escape(b["text"]),style))
            elif typ=="paragraph":
                story.append(Paragraph(escape(b["text"]),styles["body"]))
            elif typ=="bullets":
                for item in b["items"]:
                    story.append(Paragraph("• "+escape(item),styles["body"]))
            elif typ=="table":
                data=[[Paragraph(escape(str(v)),styles["th"]) for v in b["headers"]]]
                data += [[Paragraph(escape(str(v)),styles["cell"]) for v in row] for row in b["rows"]]
                widths=b["widths"] or [1]*len(b["headers"])
                widths=[doc.width*w/sum(widths) for w in widths]
                t=Table(data,colWidths=widths,repeatRows=1,hAlign="LEFT")
                t.setStyle(TableStyle([
                    ("BACKGROUND",(0,0),(-1,0),colors.HexColor("#334155")),
                    ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#f4f6f8")]),
                    ("GRID",(0,0),(-1,-1),.5,colors.HexColor("#d9d9d9")),
                    ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
                    ("LEFTPADDING",(0,0),(-1,-1),7), ("RIGHTPADDING",(0,0),(-1,-1),7),
                    ("TOPPADDING",(0,0),(-1,-1),6), ("BOTTOMPADDING",(0,0),(-1,-1),6),
                ]))
                story.extend([t,Spacer(1,9)])
            elif typ=="image":
                w,hh=image_size(b)
                story.append(KeepTogether([Image(b["path"],width=w,height=hh,hAlign="CENTER"),Paragraph(escape(b["caption"]),styles["caption"])]))
            elif typ=="code":
                story.append(Preformatted(b["text"],styles["code"],maxLineLength=85))
    doc.build(story)
    return out


def make_docx(pages):
    doc=Document()
    sec=doc.sections[0]
    sec.page_width=Inches(8.2677); sec.page_height=Inches(11.6929)
    sec.top_margin=Inches(.60); sec.bottom_margin=Inches(.60)
    sec.left_margin=Inches(.67); sec.right_margin=Inches(.67)
    sec.footer_distance=Inches(.27)
    normal=doc.styles["Normal"]
    normal.font.name="Arial"; normal.font.size=Pt(10.5)
    normal.paragraph_format.line_spacing=1.12
    normal.paragraph_format.space_after=Pt(8)
    for name,size in [("Title",24),("Heading 1",17),("Heading 2",12)]:
        s=doc.styles[name]; s.font.name="Arial"; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
        s.paragraph_format.space_before=Pt(0 if name in ("Title","Heading 1") else 9)
        s.paragraph_format.space_after=Pt(12 if name in ("Title","Heading 1") else 6)
        s.paragraph_format.keep_with_next=True
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=footer.add_run("Reinforcement Learning Lab CA   |   "); r.font.size=Pt(8)
    fld=OxmlElement("w:fldSimple"); fld.set(qn("w:instr"),"PAGE"); footer._p.append(fld)
    for pi,page in enumerate(pages):
        if pi: doc.add_page_break()
        for b in page:
            typ=b["type"]
            if typ=="heading":
                doc.add_paragraph(b["text"],style="Title" if b["level"]==0 else f"Heading {b['level']}")
            elif typ=="paragraph": doc.add_paragraph(b["text"])
            elif typ=="bullets":
                for txt in b["items"]: doc.add_paragraph(txt,style="List Bullet")
            elif typ=="table":
                t=doc.add_table(rows=1,cols=len(b["headers"])); t.autofit=False
                widths=b["widths"] or [1]*len(b["headers"])
                total=6.9277
                for j,w in enumerate(widths): t.columns[j].width=Inches(total*w/sum(widths))
                for j,txt in enumerate(b["headers"]): t.rows[0].cells[j].text=str(txt)
                for row in b["rows"]:
                    cells=t.add_row().cells
                    for j,txt in enumerate(row): cells[j].text=str(txt)
                for i,row in enumerate(t.rows):
                    trPr=row._tr.get_or_add_trPr()
                    if i==0:
                        repeat=OxmlElement("w:tblHeader"); trPr.append(repeat)
                    cant=OxmlElement("w:cantSplit"); trPr.append(cant)
                    for j,cell in enumerate(row.cells):
                        cell.width=Inches(total*widths[j]/sum(widths)); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        tcPr=cell._tc.get_or_add_tcPr()
                        shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"334155" if i==0 else "F4F6F8" if i%2==0 else "FFFFFF"); tcPr.append(shd)
                        borders=OxmlElement("w:tcBorders")
                        for edge in ["top","left","bottom","right"]:
                            tag=OxmlElement("w:"+edge); tag.set(qn("w:val"),"single"); tag.set(qn("w:sz"),"4"); tag.set(qn("w:color"),"D9D9D9"); borders.append(tag)
                        tcPr.append(borders)
                        margins=OxmlElement("w:tcMar")
                        for edge in ["top","left","bottom","right"]:
                            tag=OxmlElement("w:"+edge); tag.set(qn("w:w"),"90"); tag.set(qn("w:type"),"dxa"); margins.append(tag)
                        tcPr.append(margins)
                        for para in cell.paragraphs:
                            para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1.05
                            for run in para.runs:
                                run.font.size=Pt(9); run.font.bold=i==0; run.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
            elif typ=="image":
                w,hh=image_size(b)
                para=doc.add_paragraph(); para.alignment=WD_ALIGN_PARAGRAPH.CENTER; para.paragraph_format.space_after=Pt(0); para.paragraph_format.keep_with_next=True
                para.add_run().add_picture(b["path"],width=Inches(w/72),height=Inches(hh/72))
                para=doc.add_paragraph(b["caption"]); para.paragraph_format.space_after=Pt(8)
                for r in para.runs: r.italic=True; r.font.size=Pt(8.7)
            elif typ=="code":
                para=doc.add_paragraph(); para.paragraph_format.line_spacing=1; para.paragraph_format.space_after=Pt(9)
                run=para.add_run(b["text"]); run.font.name="Consolas"; run.font.size=Pt(8.2)
    doc.core_properties.title="Tic Tac Toe Agent Using Tabular Q Learning"
    doc.core_properties.subject="Reinforcement Learning Lab CA mini project report"
    doc.core_properties.author=""
    out=OUT/"LAB_CA_Report.docx"; doc.save(out); return out


def make_markdown(pages):
    chunks=[]
    for idx,page in enumerate(pages,1):
        for b in page:
            typ=b["type"]
            if typ=="heading": chunks.append("#"*(b["level"]+1)+" "+b["text"])
            elif typ=="paragraph": chunks.append(b["text"])
            elif typ=="bullets": chunks.append("\n".join("- "+s for s in b["items"]))
            elif typ=="table":
                chunks.append("\n".join(["| "+" | ".join(map(str,b["headers"]))+" |", "| "+" | ".join("---" for _ in b["headers"])+" |"] + ["| "+" | ".join(map(str,row))+" |" for row in b["rows"]]))
            elif typ=="image":
                rel=Path(b["path"]).relative_to(ROOT).as_posix()
                rel=rel[len("report/"):] if rel.startswith("report/") else "../"+rel
                chunks.append("!["+b["caption"]+"]("+rel+")\n\n"+b["caption"])
            elif typ=="code": chunks.append("```python\n"+b["text"]+"\n```")
        chunks.append("\n<!-- page break -->\n")
    out=OUT/"LAB_CA_Report.md"; out.write_text("\n\n".join(chunks),encoding="utf-8"); return out


def measured_context():
    summary=read_json(ART/"training_summary.json")
    evaluation=read_json(ART/"evaluation.json")
    if not summary or not evaluation:
        raise SystemExit("Run training first; measured summary and evaluation artifacts are required.")
    cfg=summary["parameters"]
    trained=evaluation["trained"]
    base=evaluation["baseline"]
    first=summary["initial_training_window"]
    last=summary["final_training_window"]
    audit=evaluation["adversarial_audit"]
    percent=lambda x:f"{x*100:.2f}"
    result_rows=[]
    for label,key in [("Baseline","baseline"),("Trained","trained")]:
        for opp in ("random","tactical","minimax"):
            row=evaluation[key][opp]["combined"]
            result_rows.append([label,opp.title(),f"{row['games']:,}",percent(row["win_rate"]),percent(row["draw_rate"]),percent(row["loss_rate"]),f"{row['average_reward']:.4f}"])
    role_rows=[]
    for opp in ("random","tactical","minimax"):
        for role in ("X","O"):
            row=trained[opp][role]
            role_rows.append([role,opp.title(),f"{row['games']:,}",f"{row['wins']:,}",f"{row['draws']:,}",f"{row['losses']:,}",percent(row["non_loss_rate"])])
    rows=list(csv.DictReader((ART/"evaluation_checkpoints.csv").open(encoding="utf-8")))
    first_checkpoint={r["opponent"]:float(r["non_loss_rate"]) for r in rows if int(r["episode"])==cfg["evaluation_interval"]}
    last_checkpoint={r["opponent"]:float(r["non_loss_rate"]) for r in rows if int(r["episode"])==cfg["episodes"]}
    agent_source=(ROOT/"tictactoe"/"agent.py").read_text(encoding="utf-8")
    update=agent_source.split("        if state[action]",1)[1].split("\n    def export",1)[0]
    update=textwrap.dedent("        if state[action]"+update).strip()
    allgames=sum(v["combined"]["games"] for v in trained.values())
    all_losses=sum(v["combined"]["losses"] for v in trained.values())
    random_win=trained["random"]["combined"]["win_rate"]
    context={
        "overview_result":f"After {cfg['episodes']:,} training episodes, the saved policy won {random_win:.2%} of {trained['random']['combined']['games']:,} test games against a random opponent and drew all {trained['minimax']['combined']['games']:,} games against minimax. It recorded {all_losses} losses across {allgames:,} final evaluation games. A separate adversarial game-tree audit found zero loss probability from the empty board for both player roles under the saved greedy policy.",
        "reward_description":"Win +1.0; draw +0.3; loss -1.0; ongoing transition 0. Terminal outcome is measured from the learner's perspective.",
        "gamma":str(cfg["gamma"]),
        "libraries":f"Python standard library for the learning engine; NumPy {summary['library_versions']['numpy']} and Matplotlib {summary['library_versions']['matplotlib']} for numerical support and plots; notebook display tools for interactive outputs.",
        "tactical_description":"Take an immediate win; otherwise block an immediate opponent win, take the center, take a corner or select a remaining edge. It does not perform full game-tree search.",
        "hyper_rows":[
            ["Learning rate α",str(cfg["alpha"])+" (constant)"],
            ["Discount factor γ",str(cfg["gamma"])],
            ["Exploration rate ε","Linear 1.00 to 0.03 during the first 136,000 episodes; 0.03 thereafter"],
            ["Number of episodes",f"{cfg['episodes']:,} (80,000 as X and 80,000 as O)"],
            ["Rewards","Win +1; draw +0.3; loss -1; ongoing 0"],
            ["Opponent mixture","Per move: random 25%, tactical 15%, minimax 60%"],
            ["Seeds","Training 42; environment 43; checkpoint 2026; final evaluation 1002026"],
            ["Logging and checkpoints","1,000-episode training windows; evaluate every 5,000 episodes"],
            ["Evaluation games","Checkpoint: 200 per role/opponent; final: 2,000 per role/opponent"],
            ["Greedy ties and initial Q","Uniform legal ties within 1e-12; unvisited values start at 0"],
        ],
        "training_description":f"Training ran locally for {cfg['episodes']:,} episodes with seed {cfg['seed']}. The learner alternated X and O on successive episodes. On each opponent move, a new policy was sampled independently from the fixed mixture of 25% random, 15% tactical and 60% minimax. The run learned {summary['canonical_states']:,} canonical decision states and exported {summary['exported_states']:,} board orientations for the browser.",
        "evaluation_protocol":f"The untrained baseline uses an all-zero Q table with uniform random selection among tied legal actions. Both baseline and final policy were tested with epsilon zero in {cfg['final_games_per_role']:,} games per role against each of three opponents: {allgames:,} games per policy. Final seed {evaluation['evaluation_seed']} differs from training and checkpoint seeds. The opponents are held fixed within each evaluation condition and Q values are not updated.",
        "screenshots":[],
        "execution_status":"The saved 160,000-episode training and final evaluation were executed locally on Windows. The self-contained browser game is opened from the local project folder. Notebook execution and interface checks are documented in the project verification artifacts.",
        "results_intro":f"Table values below come directly from artifacts/evaluation.json. Each row pools {cfg['final_games_per_role']:,} games as X and {cfg['final_games_per_role']:,} as O. The same protocol evaluates the zero-table baseline and the saved trained policy.",
        "result_rows":result_rows,
        "role_rows":role_rows,
        "results_interpretation":f"Against random play, win rate increased from {base['random']['combined']['win_rate']:.2%} to {random_win:.2%}. Against minimax, the baseline lost {base['minimax']['combined']['loss_rate']:.2%} of games, while the trained policy drew all games. The trained policy's average terminal reward was {trained['random']['combined']['average_reward']:.4f}, {trained['tactical']['combined']['average_reward']:.4f} and {trained['minimax']['combined']['average_reward']:.4f} against random, tactical and minimax opponents respectively.",
        "reward_plot":str(ART/"plots"/"reward_vs_episode.png"),
        "reward_caption":"Figure 3. Mean terminal reward during training, measured in non-overlapping windows of 1,000 episodes.",
        "reward_analysis":f"The first training window averaged {first['average_reward']:.4f} reward and the final window averaged {last['average_reward']:.4f}. These are exploratory training games against a per-move mixture, so they should not be equated with the final greedy evaluation. Total training reward was {last['cumulative_reward']:,.1f}; early exploratory losses still contribute to this cumulative total.",
        "success_plot":str(ART/"plots"/"success_vs_episode.png"),
        "success_caption":"Figure 4. Greedy checkpoint non-loss rate; 400 games per opponent at each checkpoint, using seed 2026.",
        "steps_plot":str(ART/"plots"/"steps_vs_episode.png"),
        "steps_caption":"Figure 5. Mean board moves and agent decisions during training, summarized over 1,000 episodes.",
        "curve_analysis":f"Training non-loss rate rose from {first['non_loss_rate']:.1%} to {last['non_loss_rate']:.1%}. Mean game length changed from {first['average_board_moves']:.3f} to {last['average_board_moves']:.3f} board moves, consistent with fewer early losses and more full-board draws. Longer games are not automatically worse performance in this task.",
        "analysis":f"At the first 5,000-episode checkpoint, non-loss rates were {first_checkpoint['random']:.1%} against random, {first_checkpoint['tactical']:.1%} against tactical and {first_checkpoint['minimax']:.1%} against minimax. Final checkpoint rates were {last_checkpoint['random']:.1%}, {last_checkpoint['tactical']:.1%} and {last_checkpoint['minimax']:.1%}. The policy learned to avoid defeat and exploit weak responses. A nonzero exploration floor and stochastic opponent keep training losses and temporal difference error above zero.",
        "audit_analysis":f"The exact adversarial audit explores every opponent response reachable under the saved greedy policy, averaging all maximal-Q ties. It visited {audit['X']['reachable_agent_states']} learner states as X and {audit['O']['reachable_agent_states']} as O, with no unseen learner states. Worst-case loss probability was zero for both roles. This stronger result applies from the empty board with this policy and tie rule; it is not a claim about arbitrary forced board positions or exploratory modes.",
        "conclusion":f"A tabular Q-learning agent was implemented and trained for {cfg['episodes']:,} games using legal actions, agent-perspective states, symmetry sharing and two-move transitions. The trained greedy policy won {random_win:.2%} against random play, drew every sampled minimax game and lost none of {allgames:,} final test games. An exact adversarial audit also found no losing path from the empty board for either role under the saved greedy policy. These results demonstrate effective learned play within the stated evaluation scope.",
        "provenance":f"Experiment timestamp: {summary['created_utc']}. Runtime: Python {summary['python_version']} on {summary['platform']}. The recorded training loop time was {summary['training_elapsed_seconds']:.2f} seconds; this excludes baseline evaluation, final evaluation, plotting, report creation and user interface verification.",
        "run_commands":"""# From the RL_LABCA project folder
python -m pip install -r requirements.txt
python -m unittest discover -s tests -v
python -m tictactoe.train --episodes 160000 --seed 42

# Open TicTacToe_Offline.html by double-clicking it.
# Open notebooks/TicTacToe_Q_Learning.ipynb in Jupyter.
""",
        "update_code":update,
        "sample_output_intro":"The following compact output reproduces values from the saved evaluation artifact. Outcomes are shown as wins / draws / losses for the trained policy, pooling the two player roles.",
        "sample_output":"\n".join([
            f"Training episodes: {cfg['episodes']:,}",
            f"Canonical states: {summary['canonical_states']:,}",
            f"Exported states: {summary['exported_states']:,}",
            "", "Opponent       Wins     Draws    Losses",
            *[f"{k.title():<14} {v['combined']['wins']:>5}    {v['combined']['draws']:>5}    {v['combined']['losses']:>6}" for k,v in trained.items()],
            "", f"Exact worst-case loss as X: {audit['X']['loss_probability']:.1f}",
            f"Exact worst-case loss as O: {audit['O']['loss_probability']:.1f}"]),
        "verification":["Environment tests check game rules, terminal states and legal actions.","Agent tests check terminal updates, legal bootstrapping, symmetry mapping and exported-policy behavior.","Final evaluation uses fresh seeds and disables exploration and Q-table updates.","The exact adversarial audit explores every opponent response reachable under the saved greedy policy from the empty board."],
    }
    notes=read_json(OUT/"report_notes.json",{})
    for key in ("screenshots","execution_status","verification"):
        if key in notes: context[key]=notes[key]
    return context


def main():
    OUT.mkdir(exist_ok=True)
    create_diagrams()
    context_path=OUT/"report_context.json"
    context=measured_context()
    context_path.write_text(json.dumps(context,indent=2),encoding="utf-8")
    pages=build_pages(context)
    paths=[make_pdf(pages),make_docx(pages),make_markdown(pages)]
    for path in paths: print(path)
    print(f"Content plan: {len(pages)} pages; 13 main pages, references and two appendices")


if __name__=="__main__":
    main()
